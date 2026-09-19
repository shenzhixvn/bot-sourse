import sys, os, json, shutil, subprocess

# ============================================================
#  BOT — 本地文件处理脚本
#  支持：txt/md 纯文本、docx Word、xlsx Excel
#  操作：读写、追加、复制、重命名、删除、列出、打开
# ============================================================

# ── 尝试导入 Office 处理库 ──
try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from openpyxl import Workbook, load_workbook
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


def _result(status, **kwargs):
    r = {"status": status}
    r.update(kwargs)
    return r


def _ensure_dir(file_path):
    d = os.path.dirname(file_path)
    if d and not os.path.exists(d):
        os.makedirs(d, exist_ok=True)


# ── 纯文本操作 ──

def write_file(path, content):
    _ensure_dir(path)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)
    return _result("ok", path=os.path.abspath(path), bytes=len(content.encode('utf-8')))


def read_file(path):
    with open(path, 'r', encoding='utf-8', errors='replace') as f:
        content = f.read()
    return _result("ok", content=content, path=os.path.abspath(path))


def append_file(path, content):
    _ensure_dir(path)
    with open(path, 'a', encoding='utf-8') as f:
        f.write(content)
    return _result("ok", path=os.path.abspath(path), appended=len(content.encode('utf-8')))


# ── Word (.docx) 操作 ──

def read_docx(path):
    if not HAS_DOCX:
        return _result("error", message="未安装 python-docx 库，请运行: pip install python-docx")
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    # 提取表格内容
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            tables_text.append(" | ".join(cell.text for cell in row.cells))
    content = "\n".join(paragraphs)
    if tables_text:
        content += "\n\n--- 表格 ---\n" + "\n".join(tables_text)
    return _result("ok", content=content, path=os.path.abspath(path),
                   paragraphs=len(paragraphs), tables=len(doc.tables))


def write_docx(path, content):
    if not HAS_DOCX:
        return _result("error", message="未安装 python-docx 库，请运行: pip install python-docx")
    _ensure_dir(path)
    doc = Document()
    # 按空行分段，连续非空行作为一个段落
    lines = content.split('\n')
    current_para = []
    for line in lines:
        if line.strip() == '':
            if current_para:
                doc.add_paragraph('\n'.join(current_para))
                current_para = []
        else:
            current_para.append(line)
    if current_para:
        doc.add_paragraph('\n'.join(current_para))
    doc.save(path)
    return _result("ok", path=os.path.abspath(path), paragraphs=len(doc.paragraphs))


# ── Excel (.xlsx) 操作 ──

def read_xlsx(path):
    if not HAS_XLSX:
        return _result("error", message="未安装 openpyxl 库，请运行: pip install openpyxl")
    wb = load_workbook(path, data_only=True)
    sheets = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else '' for c in row])
        sheets[sheet_name] = rows
    # 生成文本预览（第一个 sheet）
    first_sheet = wb.sheetnames[0]
    preview_rows = sheets[first_sheet]
    preview = "\n".join("\t".join(r) for r in preview_rows)
    return _result("ok", content=preview, path=os.path.abspath(path),
                   sheets=sheets, sheet_names=wb.sheetnames, active_sheet=first_sheet)


def write_xlsx(path, content):
    if not HAS_XLSX:
        return _result("error", message="未安装 openpyxl 库，请运行: pip install openpyxl")
    _ensure_dir(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # 支持两种格式：制表符分隔的行，或 CSV 风格
    lines = content.strip().split('\n')
    for line in lines:
        if '\t' in line:
            cells = line.split('\t')
        else:
            cells = [c.strip() for c in line.split(',')]
        ws.append(cells)
    wb.save(path)
    return _result("ok", path=os.path.abspath(path), rows=len(lines), cols=ws.max_column)


# ── 文件管理操作 ──

def mkdir(path):
    if not os.path.exists(path):
        os.makedirs(path, exist_ok=True)
    return _result("ok", path=os.path.abspath(path))


def list_dir(path):
    items = []
    for name in sorted(os.listdir(path)):
        full = os.path.join(path, name)
        try:
            stat = os.stat(full)
            items.append({
                "name": name,
                "size": stat.st_size if os.path.isfile(full) else None,
                "is_dir": os.path.isdir(full),
                "modified": stat.st_mtime,
            })
        except OSError:
            items.append({"name": name, "size": None, "is_dir": os.path.isdir(full), "modified": None})
    return _result("ok", items=items, path=os.path.abspath(path))


def copy_file(source, dest):
    _ensure_dir(dest)
    shutil.copy2(source, dest)
    return _result("ok", source=os.path.abspath(source), dest=os.path.abspath(dest))


def rename_file(source, dest):
    _ensure_dir(dest)
    os.rename(source, dest)
    return _result("ok", old_path=os.path.abspath(source), new_path=os.path.abspath(dest))


def delete_file(path):
    if os.path.isdir(path):
        shutil.rmtree(path)
    else:
        os.remove(path)
    return _result("ok", path=os.path.abspath(path), deleted=True)


def open_file(path):
    if sys.platform == 'win32':
        os.startfile(path)
    elif sys.platform == 'darwin':
        subprocess.call(['open', path])
    else:
        subprocess.call(['xdg-open', path])
    return _result("ok", path=os.path.abspath(path), opened=True)


def exec_code(code):
    ns = {}
    exec(code, ns)
    return _result("ok", result=str(ns.get('result', '')))


# ── 主入口 ──

ACTIONS = {
    'write': lambda a: write_file(a[2], a[3]),
    'read': lambda a: read_file(a[2]),
    'append': lambda a: append_file(a[2], a[3]),
    'mkdir': lambda a: mkdir(a[2]),
    'list': lambda a: list_dir(a[2]),
    'copy': lambda a: copy_file(a[2], a[3]),
    'rename': lambda a: rename_file(a[2], a[3]),
    'delete': lambda a: delete_file(a[2]),
    'open': lambda a: open_file(a[2]),
    'read_docx': lambda a: read_docx(a[2]),
    'write_docx': lambda a: write_docx(a[2], a[3]),
    'read_xlsx': lambda a: read_xlsx(a[2]),
    'write_xlsx': lambda a: write_xlsx(a[2], a[3]),
    'exec': lambda a: exec_code(a[2]),
}

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(json.dumps({"status": "error", "message": "缺少操作参数"}, ensure_ascii=False))
        sys.exit(1)

    action = sys.argv[1]
    try:
        if action in ACTIONS:
            r = ACTIONS[action](sys.argv)
        else:
            r = {"status": "error", "message": f"未知操作: {action}"}
        print(json.dumps(r, ensure_ascii=False))
    except Exception as e:
        print(json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False))
        sys.exit(1)
